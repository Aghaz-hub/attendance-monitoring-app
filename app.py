"""
Attendance Monitoring Report Generator
Faculty uploads ERP attendance export → selects monitoring period → downloads formatted report
"""

import streamlit as st
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from io import BytesIO
import re

# ─────────────────────────────────────────────
# Page config
# ─────────────────────────────────────────────
st.set_page_config(
    page_title="Attendance Monitoring Report Generator",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ─────────────────────────────────────────────
# Styles for generated Excel
# ─────────────────────────────────────────────
THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
HEADER_FONT = Font(bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
HEADER_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
CRITICAL_FILL = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
LIGHT_RED = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


# ─────────────────────────────────────────────
# Parser: ERP-style Attendance Monitoring Report
# ─────────────────────────────────────────────
def parse_erp_attendance(wb):
    """
    Parse the wide ERP export format (like CSTI SEM 5).
    Returns dict with metadata + list of student dicts.
    """
    ws = wb.active

    # ── Metadata ──
    meta = {
        "branch": "",
        "department": "",
        "class_name": "",
        "division": "All",
        "date_range": "",
        "academic_year": "",
        "semester": "",
    }

    for row in range(1, min(25, ws.max_row + 1)):
        for col in range(1, min(15, ws.max_column + 1)):
            val = ws.cell(row, col).value
            if val is None:
                continue
            text = str(val).strip()
            low = text.lower()

            if low.startswith("branch:") or low.startswith("branch :"):
                meta["branch"] = text.split(":", 1)[-1].strip()
            elif "department" in low and ":" in text and "dept of" not in low:
                # Prefer the longer department description
                dept = text.split(":", 1)[-1].strip()
                if len(dept) > len(meta["department"]):
                    meta["department"] = dept
            elif low.startswith("class name:") or (
                low.startswith("class :") and "cse" in low
            ):
                meta["class_name"] = text.split(":", 1)[-1].strip()
            elif low.startswith("division:") and "section" not in low:
                meta["division"] = text.split(":", 1)[-1].strip()
            elif "from date" in low or (low.startswith("date:") and "to" in low):
                meta["date_range"] = text.replace("Date:", "").replace("date:", "").strip()
            elif "academic year" in low:
                meta["academic_year"] = text.split(":", 1)[-1].strip()
            elif low.startswith("semester"):
                meta["semester"] = text.split(":", 1)[-1].strip()

    if not meta["branch"]:
        meta["branch"] = "MRU-School of Engineering"
    if not meta["division"] or meta["division"].lower() in ("division/section",):
        meta["division"] = "All"

    # ── Locate header row (subject names + student identity columns) ──
    subject_row = None
    roll_col = name_col = None

    for row in range(1, min(30, ws.max_row + 1)):
        vals = [str(ws.cell(row, c).value or "").strip().lower() for c in range(1, 12)]
        joined = " ".join(vals)
        if "student name" in joined and ("roll" in joined or "sr" in joined):
            subject_row = row
            for c in range(1, min(12, ws.max_column + 1)):
                hdr = str(ws.cell(row, c).value or "").strip().lower()
                if hdr in ("rollno", "roll no", "roll no.", "roll number") or (
                    "roll" in hdr and "enroll" not in hdr and "prn" not in hdr
                ):
                    roll_col = c
                if "student name" in hdr or hdr == "name":
                    name_col = c
            break

    if subject_row is None:
        raise ValueError(
            "Could not find 'Student Name' / subject header row. "
            "Please upload a standard ERP Attendance Monitoring Report."
        )

    # Detect structure under subject header
    # Typical ERP:
    #   subject names → PP/PR/Tut component → codes → CORE/ELECTIVE → D/A/% headers → faculty → data
    next_sample = str(
        ws.cell(subject_row + 1, 7).value or ws.cell(subject_row + 1, 4).value or ""
    ).strip().upper()
    if next_sample in ("PP", "PR", "T", "P", "TH", "LAB", "TUT", "TUTORIAL"):
        component_row = subject_row + 1   # PP / PR / Tut
        code_row = subject_row + 2
        type_row = subject_row + 3
        pct_header_row = subject_row + 4  # D (PP) | A (PP) | % (PP) | ...
        faculty_row = subject_row + 5
    else:
        component_row = None
        code_row = subject_row + 1
        type_row = subject_row + 2
        pct_header_row = subject_row + 3
        faculty_row = subject_row + 4

    data_start = faculty_row + 1
    while data_start <= ws.max_row:
        sr = ws.cell(data_start, 1).value
        if sr is not None and str(sr).strip().replace(".", "").isdigit():
            break
        data_start += 1

    if roll_col is None:
        roll_col = 4
    if name_col is None:
        name_col = 5

    # ── Build subject list ──
    # Rules:
    #  - %(PP) / Theory  → separate Theory column
    #  - %(PR) / Practical → separate Practical column
    #  - %(Tut) / Tutorial → counted under Theory only
    #  - If one subject name spans both PP and PR → split into Theory + Practical
    SKIP = {
        "sr.", "sr", "sr no.", "sr no", "division/section", "division",
        "unique id", "uniqueid", "rollno", "roll no", "roll no.", "roll number",
        "student name", "name", "prn / enroll", "prn", "prn/enroll", "total",
        "enroll", "enrollment", "total d", "total a", "total %",
    }

    def _component_kind(text):
        """Classify a header cell as PP (theory), PR (practical), Tut, or unknown."""
        t = (text or "").strip().upper()
        if not t:
            return None
        if "TUT" in t or "TUTORIAL" in t:
            return "TUT"
        if "PR" in t and "PP" not in t:
            return "PR"
        if "PP" in t or t in ("T", "TH", "THEORY"):
            return "PP"
        if t in ("P", "LAB", "PRACTICAL"):
            return "PR"
        return None

    def _display_name(base_name, kind, already_lab_in_name):
        """Build output column title: Theory vs Practical."""
        base = base_name.strip()
        if kind in ("PP", "TUT"):
            return base
        if kind == "PR":
            if already_lab_in_name or "LAB" in base.upper() or "PRACTICAL" in base.upper():
                return base
            # Project / practical without LAB in name
            if "PROJECT" in base.upper() or "INTERNSHIP" in base.upper():
                return base
            return f"{base} LAB"
        return base

    # Map each % column → subject entry
    # First, find all subject name positions on subject_row
    subject_positions = []  # (col, name)
    total_header_col = None
    for c in range(1, ws.max_column + 1):
        raw = ws.cell(subject_row, c).value
        name = str(raw).strip() if raw else ""
        if not name:
            continue
        low = name.lower()
        if low in ("total", "total %", "overall"):
            total_header_col = c
            continue
        if low not in SKIP:
            subject_positions.append((c, name))

    subjects = []
    used_pct_cols = set()

    for idx, (start_col, base_name) in enumerate(subject_positions):
        # Column range for this subject: until next subject name or Total column
        if idx + 1 < len(subject_positions):
            end_col = subject_positions[idx + 1][0]
        elif total_header_col:
            end_col = total_header_col
        else:
            end_col = min(start_col + 9, ws.max_column + 1)

        code = str(ws.cell(code_row, start_col).value or "").strip()
        typ = str(ws.cell(type_row, start_col).value or "CORE").strip() or "CORE"
        faculty = str(ws.cell(faculty_row, start_col).value or "").strip()
        already_lab = any(
            k in base_name.upper() for k in ("LAB", "PRACTICAL", "PRACTICUM")
        )

        # Collect component blocks under this subject: look at pct_header_row
        # Pattern is groups of D / A / % with optional (PP)/(PR)/(Tut) labels
        components = []  # list of {kind, pct_col, d_col, faculty}
        c = start_col
        while c < end_col:
            hdr = str(ws.cell(pct_header_row, c).value or "").strip()
            hdr_l = hdr.lower()
            kind = _component_kind(hdr)

            # Also check component_row (PP/PR row above codes) for this column
            if kind is None and component_row:
                kind = _component_kind(str(ws.cell(component_row, c).value or ""))

            if "%" in hdr_l or "percent" in hdr_l:
                # This is a % column — find matching D column (usually 2 left)
                d_col = c - 2 if c - 2 >= start_col else c
                # Faculty may be aligned to the D column of this component
                fac = str(ws.cell(faculty_row, d_col).value or "").strip() or faculty
                if kind is None:
                    # Fallback: use component_row at d_col / start
                    if component_row:
                        kind = _component_kind(
                            str(ws.cell(component_row, d_col).value or "")
                        ) or _component_kind(
                            str(ws.cell(component_row, start_col).value or "")
                        )
                if kind is None:
                    kind = "PR" if already_lab else "PP"
                components.append({
                    "kind": kind,
                    "pct_col": c,
                    "d_col": d_col,
                    "faculty": fac,
                })
                used_pct_cols.add(c)
                c += 1
            elif hdr_l.startswith("d ") or hdr_l.startswith("d(") or hdr_l == "d":
                # Start of a D/A/% triplet — peek +2 for %
                pct_c = c + 2
                if pct_c < end_col and pct_c not in used_pct_cols:
                    pct_hdr = str(ws.cell(pct_header_row, pct_c).value or "").lower()
                    if "%" in pct_hdr or "percent" in pct_hdr or not pct_hdr:
                        k = kind or _component_kind(
                            str(ws.cell(pct_header_row, pct_c).value or "")
                        )
                        if k is None and component_row:
                            k = _component_kind(
                                str(ws.cell(component_row, c).value or "")
                            )
                        if k is None:
                            k = "PR" if already_lab else "PP"
                        fac = str(ws.cell(faculty_row, c).value or "").strip() or faculty
                        components.append({
                            "kind": k,
                            "pct_col": pct_c,
                            "d_col": c,
                            "faculty": fac,
                        })
                        used_pct_cols.add(pct_c)
                        c = pct_c + 1
                        continue
                c += 1
            else:
                c += 1

        # If no components found via headers, fall back to classic 3-col block
        if not components:
            pct_col = start_col + 2
            kind = None
            if component_row:
                kind = _component_kind(str(ws.cell(component_row, start_col).value or ""))
            if kind is None:
                kind = "PR" if already_lab else "PP"
            components = [{
                "kind": kind,
                "pct_col": pct_col,
                "d_col": start_col,
                "faculty": faculty,
            }]

        # Merge Tut into Theory (PP): if both PP and Tut exist, prefer PP % for theory;
        # if only Tut, treat as Theory. Always emit separate PR when present.
        theory_comp = None
        practical_comp = None
        for comp in components:
            if comp["kind"] in ("PP", "TUT"):
                # Prefer PP over Tut when both exist
                if theory_comp is None or (
                    theory_comp["kind"] == "TUT" and comp["kind"] == "PP"
                ):
                    theory_comp = comp
            elif comp["kind"] == "PR":
                practical_comp = comp

        if theory_comp:
            disp = _display_name(base_name, "PP", already_lab)
            # When both theory and practical exist under same name, tag Theory clearly
            if practical_comp and "LAB" not in disp.upper():
                disp_theory = f"{base_name} (Theory)"
            else:
                disp_theory = disp
            subjects.append({
                "name": disp_theory,
                "code": code,
                "type": typ,
                "pct_col": theory_comp["pct_col"],
                "d_col": theory_comp["d_col"],
                "faculty": theory_comp["faculty"],
                "component": "Theory",
            })

        if practical_comp:
            disp_pr = _display_name(base_name, "PR", already_lab)
            if theory_comp and "LAB" not in disp_pr.upper() and "(Practical)" not in disp_pr:
                # If theory already took the base name, make practical distinct
                if not already_lab:
                    disp_pr = f"{base_name} (Practical)"
            subjects.append({
                "name": disp_pr,
                "code": code + ("-PR" if theory_comp else ""),
                "type": typ,
                "pct_col": practical_comp["pct_col"],
                "d_col": practical_comp["d_col"],
                "faculty": practical_comp["faculty"],
                "component": "Practical",
            })

        # Edge case: components found but none classified — emit first as-is
        if not theory_comp and not practical_comp and components:
            comp = components[0]
            subjects.append({
                "name": base_name,
                "code": code,
                "type": typ,
                "pct_col": comp["pct_col"],
                "d_col": comp["d_col"],
                "faculty": comp["faculty"],
                "component": comp["kind"],
            })

    # Total % column (prefer explicit Total header; avoid subject % cols)
    total_pct_col = None
    for c in range(1, ws.max_column + 1):
        v = str(ws.cell(subject_row, c).value or "").strip().lower()
        if v in ("total", "total %", "overall"):
            # Total block is usually D, A, % → % at c+2
            for off in (2, 1, 0):
                if c + off <= ws.max_column:
                    total_pct_col = c + off
                    break
            break
    if total_pct_col is None:
        # last % column that was NOT used as a subject %
        for c in range(ws.max_column, 0, -1):
            h = str(ws.cell(pct_header_row, c).value or "").lower()
            if ("%" in h or "percent" in h) and c not in used_pct_cols:
                total_pct_col = c
                break

    # ── Parse students ──
    students = []
    for row in range(data_start, ws.max_row + 1):
        sr = ws.cell(row, 1).value
        if sr is None:
            continue

        roll = str(ws.cell(row, roll_col).value or "").strip()
        name = str(ws.cell(row, name_col).value or "").strip()
        if not name or not roll:
            continue

        pcts = {}
        for s in subjects:
            key = s["code"] or s["name"]
            raw = ws.cell(row, s["pct_col"]).value
            d_raw = ws.cell(row, s["d_col"]).value
            try:
                d_val = float(d_raw) if d_raw not in (None, "") else 0
            except (ValueError, TypeError):
                d_val = 0
            try:
                p = float(raw) if raw not in (None, "") else None
            except (ValueError, TypeError):
                p = None
            pcts[key] = {"pct": p, "d": d_val}

        total_pct = None
        if total_pct_col:
            try:
                total_pct = float(ws.cell(row, total_pct_col).value)
            except (ValueError, TypeError):
                pass

        students.append({
            "roll": roll,
            "name": name,
            "pcts": pcts,
            "total_pct": total_pct,
        })

    return {
        "meta": meta,
        "subjects": subjects,
        "students": students,
    }



# ─────────────────────────────────────────────
# Report generator
# ─────────────────────────────────────────────
def generate_report(parsed, monitoring_label, min_criteria=75, program_coordinator=""):
    """
    Build Excel workbook matching the sample Attendance Monitoring format.
    monitoring_label: "FIRST" | "SECOND" | "THIRD"
    """
    meta = parsed["meta"]
    subjects = parsed["subjects"]
    students = parsed["students"]
    n_subj = len(subjects)

    wb = Workbook()
    ws = wb.active
    ws.title = f"{monitoring_label.title()} Att Monitoring"

    overall_col = 4 + n_subj
    roll_dup_col = overall_col + 1
    count_col = overall_col + 2
    crit_col = overall_col + 3

    # ── Title ──
    semester_hint = meta.get("semester") or meta.get("academic_year") or ""
    title = (
        f"DEPARTMENT OF CST\n"
        f"{semester_hint}\n"
        f"{monitoring_label.upper()} ATTENDANCE MONITORING REPORT"
    )
    last_col_letter = get_column_letter(crit_col)
    ws.merge_cells(f"A2:{last_col_letter}2")
    ws["A2"] = title
    ws["A2"].font = TITLE_FONT
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 55

    # ── Meta block ──
    ws["A5"] = f"Branch: {meta.get('branch') or 'MRU-School of Engineering'}"
    ws["A6"] = f"Department: {meta.get('department') or ''}"
    ws["A7"] = f"Class Name: {meta.get('class_name') or ''}"
    ws["A8"] = f"Division: {meta.get('division') or 'All'}"
    ws["A9"] = f"Date: {meta.get('date_range') or ''}"
    ws["A10"] = f"Program Coordinator: {program_coordinator}"
    for r in range(5, 11):
        ws[f"A{r}"].font = Font(bold=True, size=10)

    # ── Header row 12: subject names ──
    ws["A12"] = "Sr No."
    ws["B12"] = "Roll No"
    ws["C12"] = "Student Name"
    for i, s in enumerate(subjects):
        cell = ws.cell(12, 4 + i, s["name"])
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.fill = HEADER_FILL
        cell.border = THIN

    ws.cell(12, count_col).value = (
        "Please update the Minimum Attendance Criteria "
        "(Engineering, Management & Sciences(75%), Law (70%), Education (80%)) "
        "in the cell below"
    )
    ws.cell(12, count_col).alignment = CENTER
    ws.cell(12, count_col).font = Font(size=8, bold=True)

    for col in (1, 2, 3):
        c = ws.cell(12, col)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = THIN
        c.alignment = CENTER

    # ── Row 13: codes + min criteria ──
    for i, s in enumerate(subjects):
        cell = ws.cell(13, 4 + i, s["code"])
        cell.font = Font(bold=True, size=9)
        cell.alignment = CENTER
        cell.border = THIN

    cell = ws.cell(13, count_col, min_criteria)
    cell.font = Font(bold=True, size=12, color="FF0000")
    cell.alignment = CENTER
    cell.fill = YELLOW_FILL
    cell.border = THIN

    # ── Row 14: type + overall headers ──
    for i, s in enumerate(subjects):
        cell = ws.cell(14, 4 + i, s["type"])
        cell.font = Font(size=8)
        cell.alignment = CENTER
        cell.border = THIN

    for col, label in (
        (overall_col, "Overall %age of all subjects from ERP report"),
        (roll_dup_col, "Roll No"),
        (count_col, "Count of Courses with attendance below minimum attendance criteria"),
        (crit_col, "Whether Critical"),
    ):
        cell = ws.cell(14, col, label)
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.fill = HEADER_FILL
        cell.border = THIN

    # ── Row 15: faculty ──
    ws["A15"] = "Faculty Member Name"
    ws["A15"].font = Font(bold=True, italic=True, size=9)
    for i, s in enumerate(subjects):
        cell = ws.cell(15, 4 + i, s.get("faculty") or "")
        cell.font = Font(size=7, italic=True)
        cell.alignment = CENTER
        cell.border = THIN

    # ── Data rows ──
    data_start = 16
    subj_enrolled = [0] * n_subj
    subj_below = {t: [0] * n_subj for t in (75, 70, 65, 60)}

    for idx, stu in enumerate(students):
        row = data_start + idx
        ws.cell(row, 1, idx + 1).border = THIN
        ws.cell(row, 1).alignment = CENTER
        ws.cell(row, 2, stu["roll"]).border = THIN
        ws.cell(row, 2).alignment = CENTER
        ws.cell(row, 3, stu["name"]).border = THIN
        ws.cell(row, 3).alignment = LEFT

        below_count = 0
        for i, s in enumerate(subjects):
            key = s["code"] or s["name"]
            info = stu["pcts"].get(key, {})
            pct = info.get("pct")
            d_val = info.get("d", 0)
            col = 4 + i

            if d_val and d_val > 0:
                subj_enrolled[i] += 1
                if pct is not None:
                    val = round(pct, 2) if pct != int(pct) else int(pct)
                    cell = ws.cell(row, col, val)
                    if pct < min_criteria:
                        below_count += 1
                        cell.fill = LIGHT_RED
                    for thresh in (75, 70, 65, 60):
                        if pct < thresh:
                            subj_below[thresh][i] += 1
                else:
                    ws.cell(row, col, None)
            else:
                ws.cell(row, col, None)

            ws.cell(row, col).border = THIN
            ws.cell(row, col).alignment = CENTER

        # Overall
        tp = stu.get("total_pct")
        cell = ws.cell(row, overall_col, round(tp, 2) if tp is not None else None)
        cell.border = THIN
        cell.alignment = CENTER
        cell.font = Font(bold=True)

        # Roll dup
        cell = ws.cell(row, roll_dup_col, stu["roll"])
        cell.border = THIN
        cell.alignment = CENTER

        # Count below
        cell = ws.cell(row, count_col, below_count)
        cell.border = THIN
        cell.alignment = CENTER
        if below_count > 0:
            cell.fill = LIGHT_RED

        # Critical
        is_crit = "CRITICAL" if below_count > 0 else ""
        cell = ws.cell(row, crit_col, is_crit)
        cell.border = THIN
        cell.alignment = CENTER
        if is_crit:
            cell.fill = CRITICAL_FILL
            cell.font = Font(bold=True, color="FFFFFF")

    last_data = data_start + len(students) - 1

    # ── Summary block ──
    sum_row = last_data + 2
    labels = [
        ("Number of Students in course", subj_enrolled),
        ("Number of students below 75%", subj_below[75]),
        ("Number of students below 70%", subj_below[70]),
        ("Number of students below 65%", subj_below[65]),
        ("Number of students below 60%", subj_below[60]),
    ]
    for offset, (label, data) in enumerate(labels):
        r = sum_row + offset
        ws.cell(r, 2, label).font = Font(bold=True)
        for i, v in enumerate(data):
            cell = ws.cell(r, 4 + i, v)
            cell.alignment = CENTER
            cell.border = THIN
            if offset == 1 and v > 0:
                cell.fill = LIGHT_RED

    # %age below 75%
    r = sum_row + 5
    ws.cell(r, 2, "%age of students below 75% -").font = Font(bold=True)
    for i in range(n_subj):
        pct_b = round(subj_below[75][i] / subj_enrolled[i] * 100) if subj_enrolled[i] else 0
        cell = ws.cell(r, 4 + i, pct_b)
        cell.alignment = CENTER
        cell.border = THIN
        cell.font = Font(bold=True)

    ws.cell(sum_row + 6, 2, f"{monitoring_label.title()} Att Monitoring").font = Font(
        bold=True, italic=True
    )

    # ── Observations ──
    obs = sum_row + 9
    for i, label in enumerate(
        ("Observations:", "Action Taken:", "Action Taken for Critical Cases:")
    ):
        ws.cell(obs + i, 1, label).font = Font(bold=True, size=11)

    # ── Signatures ──
    sig = obs + 6
    for i, label in enumerate(
        (
            "Signature of Mentor",
            "Mentor Name:",
            "",
            "HoD",
            "",
            "Dean of School",
            "",
            "Associate Dean (Academics)",
            "",
            "Dean (Academics)",
        )
    ):
        if label:
            ws.cell(sig + i, max(12, n_subj // 2 + 2), label).font = Font(
                bold=True if label not in ("Mentor Name:",) else False
            )

    # ── Column widths ──
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 36
    for i in range(n_subj):
        ws.column_dimensions[get_column_letter(4 + i)].width = 11
    ws.column_dimensions[get_column_letter(overall_col)].width = 12
    ws.column_dimensions[get_column_letter(roll_dup_col)].width = 14
    ws.column_dimensions[get_column_letter(count_col)].width = 12
    ws.column_dimensions[get_column_letter(crit_col)].width = 12

    ws.row_dimensions[12].height = 48
    ws.row_dimensions[13].height = 20
    ws.row_dimensions[14].height = 32
    ws.row_dimensions[15].height = 22
    ws.freeze_panes = "D16"

    return wb


# ─────────────────────────────────────────────
# Streamlit UI
# ─────────────────────────────────────────────
def main():
    st.markdown(
        """
        <style>
        .main-title {
            font-size: 1.8rem;
            font-weight: 700;
            color: #1a365d;
            margin-bottom: 0.2rem;
        }
        .sub-title {
            color: #4a5568;
            font-size: 0.95rem;
            margin-bottom: 1.5rem;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )

    st.markdown(
        '<p class="main-title">📊 Attendance Monitoring Report Generator</p>',
        unsafe_allow_html=True,
    )
    st.markdown(
        '<p class="sub-title">Upload ERP attendance export → choose monitoring period → download formatted report</p>',
        unsafe_allow_html=True,
    )

    # ── Sidebar ──
    with st.sidebar:
        st.header("⚙️ Settings")
        monitoring = st.radio(
            "Monitoring Period",
            options=["FIRST", "SECOND", "THIRD"],
            index=1,
            help="Select which attendance monitoring cycle this report is for.",
        )
        min_criteria = st.number_input(
            "Minimum Attendance Criteria (%)",
            min_value=50,
            max_value=100,
            value=75,
            step=1,
            help="Engineering / Management / Sciences default = 75%. Law = 70%, Education = 80%.",
        )
        coordinator = st.text_input(
            "Program Coordinator",
            value="",
            placeholder="e.g. Mr. / Ms. Name",
        )
        st.divider()
        st.caption("Supported input: ERP Attendance Monitoring Report (.xlsx)")
        st.caption("Output format matches the standard departmental template.")

    # ── Upload ──
    uploaded = st.file_uploader(
        "Upload Attendance Monitoring Excel file",
        type=["xlsx", "xls"],
        help="Upload the ERP-exported Attendance Monitoring Report (wide format with D / A / % per subject).",
    )

    if uploaded is None:
        st.info(
            "👆 Upload an ERP attendance file to begin. "
            "You can use the CSTI Sem 5 file as a test."
        )
        with st.expander("Expected input format"):
            st.markdown(
                """
                The uploaded file should be an **ERP Attendance Monitoring Report** containing:
                - Header block (Branch, Department, Class, Date range)
                - Subject columns with **D (Delivered)**, **A (Attended)**, **% (Percentage)** for each course
                - Student rows with Roll No, Name, and per-subject attendance
                - A **Total** column with overall attendance %

                The app will transform it into the standard departmental monitoring format with:
                - Only % values per subject
                - Overall %
                - Count of courses below criteria
                - **CRITICAL** flag for students with any shortfall
                - Summary statistics & signature block
                """
            )
        return

    # ── Process ──
    try:
        wb_in = openpyxl.load_workbook(BytesIO(uploaded.read()), data_only=True)
    except Exception as e:
        st.error(f"Could not read the Excel file: {e}")
        return

    with st.spinner("Parsing attendance data…"):
        try:
            parsed = parse_erp_attendance(wb_in)
        except Exception as e:
            st.error(f"Parse error: {e}")
            st.exception(e)
            return

    n_stu = len(parsed["students"])
    n_subj = len(parsed["subjects"])
    meta = parsed["meta"]

    if n_stu == 0:
        st.warning("No student records found. Please check the file format.")
        return

    # ── Preview ──
    col1, col2, col3, col4 = st.columns(4)
    col1.metric("Students", n_stu)
    col2.metric("Subjects", n_subj)
    col3.metric("Criteria", f"{min_criteria}%")
    col4.metric("Monitoring", monitoring)

    st.markdown("#### Detected Class Info")
    m1, m2 = st.columns(2)
    with m1:
        st.write(f"**Branch:** {meta.get('branch') or '—'}")
        st.write(f"**Department:** {meta.get('department') or '—'}")
        st.write(f"**Class:** {meta.get('class_name') or '—'}")
    with m2:
        st.write(f"**Division:** {meta.get('division') or '—'}")
        st.write(f"**Date range:** {meta.get('date_range') or '—'}")
        st.write(f"**Academic Year:** {meta.get('academic_year') or '—'}")

    with st.expander("Subjects detected"):
        for s in parsed["subjects"]:
            st.write(
                f"- **{s['name']}** (`{s['code']}`) — {s['type']} — Faculty: {s.get('faculty') or '—'}"
            )

    # Quick critical preview
    critical_preview = []
    for stu in parsed["students"]:
        below = 0
        for s in parsed["subjects"]:
            key = s["code"] or s["name"]
            info = stu["pcts"].get(key, {})
            if (
                info.get("d", 0) > 0
                and info.get("pct") is not None
                and info["pct"] < min_criteria
            ):
                below += 1
        if below > 0:
            critical_preview.append(
                {
                    "Roll No": stu["roll"],
                    "Name": stu["name"],
                    "Courses below criteria": below,
                    "Overall %": stu.get("total_pct"),
                }
            )

    st.markdown(f"#### Critical Students Preview ({len(critical_preview)} / {n_stu})")
    if critical_preview:
        st.dataframe(critical_preview, use_container_width=True, hide_index=True)
    else:
        st.success("No students currently below the minimum attendance criteria.")

    # ── Generate & Download ──
    st.divider()
    if st.button("🚀 Generate Formatted Report", type="primary", use_container_width=True):
        with st.spinner("Building report…"):
            wb_out = generate_report(
                parsed,
                monitoring_label=monitoring,
                min_criteria=min_criteria,
                program_coordinator=coordinator,
            )
            buf = BytesIO()
            wb_out.save(buf)
            buf.seek(0)

        class_slug = re.sub(r"[^\w]+", "_", meta.get("class_name") or "Attendance")
        fname = f"{class_slug}_{monitoring.title()}_Att_Monitoring.xlsx"

        st.success(
            f"✅ Report generated — {n_stu} students, {len(critical_preview)} critical."
        )
        st.download_button(
            label="⬇️ Download Report",
            data=buf.getvalue(),
            file_name=fname,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )


if __name__ == "__main__":
    main()
